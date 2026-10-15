"""gendive-skills-common 실습용 가상 데이터 생성 (난수 seed 고정).

실행: python3 samples/make_samples.py
출력: repos/gendive-skills-common/samples/
  - 고객_방문이력_원본.xlsx  (마스킹 실습)
  - 매입내역_정리전.xlsx      (엑셀 정리 실습)
  - 작업일지_메모.txt         (메모를 표로 실습)
  - README.md                 (함정 목록 + 검산용 정답, 이 스크립트가 계산해서 기록)
모든 이름·번호·회사는 가상이며 실존 인물·업체와 무관하다.
"""
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

OUT = Path(__file__).resolve().parent
SEED = 20261015


# ---------------------------------------------------------------- 1. 고객 방문이력
def make_customers(rng):
    surnames = ['김', '이', '박', '최', '정', '강', '조', '윤', '장', '임', '한', '오', '서', '신', '권', '황', '안', '송', '류', '홍', '남궁', '선우']
    given = ['민', '서', '준', '현', '지', '우', '하', '은', '수', '영', '진', '아', '도', '예', '시', '유', '태', '린', '재', '빈']
    cities = [('서울시', ['가상구', '샘플구', '예시구']), ('경기도 수원시', ['가상구', '예시구']),
              ('부산시', ['샘플구', '테스트구']), ('인천시', ['가상구', '예시구'])]
    roads = ['샘플로', '가상대로', '예시길', '테스트로', '더미길']
    models = ['경차', '소형 세단', '준중형 세단', '중형 세단', '중형 SUV', '대형 SUV', '소형 트럭', '승합차']
    regions = ['서울', '경기', '부산', '인천']
    hangul = ['가', '나', '다', '라', '마', '거', '너', '더', '러', '머', '고', '노', '도', '로', '모']
    rows, used_phone = [], set()
    base = date(2026, 10, 14)
    for i in range(150):
        name = rng.choice(surnames) + rng.choice(given) + rng.choice(given)
        while True:
            p = rng.randint(1, 9999)
            if p not in used_phone:
                used_phone.add(p)
                break
        phone = f'010-0000-{p:04d}'
        yy, mm, dd = rng.randint(55, 99), rng.randint(1, 12), rng.randint(1, 28)
        rrn = f'{yy:02d}{mm:02d}{dd:02d}-{rng.choice([1, 2])}000000'  # 뒤 6자리 0 채움 = 실제로 유효하지 않은 형태
        city, gus = rng.choice(cities)
        addr = f'{city} {rng.choice(gus)} {rng.choice(roads)} {rng.randint(1, 150)}'
        email = f'customer{i + 1:03d}@example.com'
        k = rng.random()
        if k < 0.4:
            plate = f'{rng.randint(10, 99)}{rng.choice(hangul)}{rng.randint(1000, 9999)}'
        elif k < 0.8:
            plate = f'{rng.randint(100, 999)}{rng.choice(hangul)}{rng.randint(1000, 9999)}'
        else:
            plate = f'{rng.choice(regions)}{rng.randint(10, 99)}{rng.choice(hangul)}{rng.randint(1000, 9999)}'
        last = base - timedelta(days=rng.randint(1, 900))
        total = rng.randint(30, 4000) * 1000 + rng.choice([0, 0, 500, 450])
        rows.append([name, phone, rrn, addr, email, plate, rng.choice(models), last, total])
    return rows


def write_customers(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = '방문이력'
    ws.append(['고객명', '휴대폰', '주민등록번호', '주소', '이메일', '차량번호', '차종', '최근방문일', '누적금액(원)'])
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=8):
        row[0].number_format = 'yyyy-mm-dd'
    for row in ws.iter_rows(min_row=2, min_col=9, max_col=9):
        row[0].number_format = '#,##0'
    for col, w in zip('ABCDEFGHI', [10, 15, 17, 28, 28, 13, 14, 13, 14]):
        ws.column_dimensions[col].width = w
    wb.save(OUT / '고객_방문이력_원본.xlsx')


# ---------------------------------------------------------------- 2. 매입내역
# (표준 품목명, 규격, 단위, 공급 단가(원, 10의 배수), 표기 변형들)
ITEMS = [
    ('볼트 M8', 'M8x30', 'EA', 60, ['볼트M8', '볼트 M8', 'BOLT-M8', '볼트 Ｍ８']),
    ('너트 M8', 'M8', 'EA', 30, ['너트M8', '너트 M8', 'NUT-M8']),
    ('와셔 M8', 'M8', 'EA', 10, ['와셔M8', '와셔 M8', 'WASHER-M8']),
    ('베어링 6204', '6204ZZ', 'EA', 4500, ['베어링6204', '베어링 6204', 'BEARING 6204']),
    ('윤활유 20L', '20L', '통', 85000, ['윤활유20L', '윤활유 20L', '윤활유(20L)']),
    ('절삭유 18L', '18L', '통', 62000, ['절삭유18L', '절삭유 18L', '절삭유(18L)']),
    ('포장박스 중', '400x300x200', 'EA', 1200, ['포장박스(중)', '포장박스 중', '박스 중형']),
    ('스테인리스판 1.5T', '1.5Tx1000x2000', '장', 48000, ['스테인리스판1.5T', 'STS판 1.5T', '스테인리스판 1.5T']),
    ('코팅장갑', 'L', '켤레', 1500, ['코팅장갑', '코팅 장갑', '장갑(코팅)']),
    ('용접봉 2.6mm', '2.6mm 5kg', 'BOX', 32000, ['용접봉2.6', '용접봉 2.6mm', '용접봉 2.6㎜']),
]
VENDORS = ['가상산업(주)', '샘플철강', '예시유통', '테스트물산', '(주)더미부품']
WS_ = '　'  # 전각 공백


def fmt_date(rng, d):
    k = rng.choice(['dt', 'dt', 'dot', 'slash', 'ymd', 'serial', 'kor'])
    if k == 'dt':
        return d
    if k == 'dot':
        return d.strftime('%Y.%m.%d')
    if k == 'slash':
        return f'{d.year % 100}/{d.month}/{d.day}'
    if k == 'ymd':
        return d.strftime('%Y%m%d')
    if k == 'serial':
        return (d - date(1899, 12, 30)).days
    return f'{d.month}월 {d.day}일'


def make_purchases(rng):
    """반환: 화면에 쓸 월별 행 블록(원본 표기) + 정답용 정리 행 목록."""
    months = [(2026, 7), (2026, 8), (2026, 9)]
    per_month = 80
    truth, blocks = [], []
    for (y, m) in months:
        days = sorted(rng.randint(1, 28) for _ in range(per_month))
        block = []
        for day in days:
            name, spec, unit, up, variants = rng.choice(ITEMS)
            vendor = rng.choice(VENDORS) if rng.random() > 0.03 else ''  # 일부 거래처 빈칸
            qty = rng.choice([10, 20, 50, 100, 200, 300, 500, 1000]) if up < 5000 else rng.choice([1, 2, 3, 5, 8, 10, 20])
            supply = qty * up
            vat_mode = rng.choice(['별도', '별도', '포함'])
            thousand = vat_mode == '별도' and up % 1000 == 0 and rng.random() < 0.45
            d = date(y, m, day)
            truth.append(dict(date=d, vendor=vendor, item=name, supply=supply))
            unit_price_disp = up * 11 // 10 if vat_mode == '포함' else up  # 포함이면 단가도 부가세 포함가
            amount_disp = qty * unit_price_disp
            amt_unit = '원'
            if thousand:
                unit_price_disp //= 1000
                amount_disp //= 1000
                amt_unit = '천원'
            name_disp = rng.choice(variants) if rng.random() < 0.75 else name
            if rng.random() < 0.12:
                name_disp = ' ' + name_disp + ' '
            if rng.random() < 0.05:
                name_disp = name_disp.replace(' ', WS_)
            qty_disp = qty
            r = rng.random()
            if r < 0.15:
                qty_disp = f'{qty:,}'
            elif r < 0.22:
                qty_disp = f'{qty}{"개" if unit == "EA" else unit}'
            elif r < 0.26:
                qty_disp = f' {qty} '
            up_disp = unit_price_disp
            if rng.random() < 0.12:
                up_disp = f'₩{unit_price_disp:,}'
            elif rng.random() < 0.08:
                up_disp = f'{unit_price_disp:,}원'
            amt_disp = amount_disp
            if rng.random() < 0.12:
                amt_disp = f'{amount_disp:,}'
            vendor_disp = vendor
            if vendor and rng.random() < 0.1:
                vendor_disp = vendor + ' '
            block.append([fmt_date(rng, d), vendor_disp, name_disp, spec, unit, qty_disp, up_disp, amt_disp, vat_mode, amt_unit, None])
        blocks.append(block)
    return blocks, truth


def write_purchases(blocks, rng):
    wb = Workbook()
    ws = wb.active
    ws.title = '매입내역'
    ws.merge_cells('A1:K1')
    ws['A1'] = '2026년 3분기 매입 내역 (자재팀)'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = Alignment(horizontal='center')
    ws.merge_cells('A2:E2')
    ws['A2'] = '거래 정보'
    ws.merge_cells('F2:J2')
    ws['F2'] = '수량·금액 정보'
    ws.merge_cells('K2:K3')
    ws['K2'] = '비고'
    for c in ('A2', 'F2', 'K2'):
        ws[c].alignment = Alignment(horizontal='center')
        ws[c].font = Font(bold=True)
    ws.append(['일자', '거래처', '품목', '규격', '단위', '수량', '단가', '금액', '부가세구분', '금액단위', None])
    for c in ws[3]:
        c.font = Font(bold=True)
    month_names = ['7월', '8월', '9월']
    grand = 0
    data_rows = []
    for mi, block in enumerate(blocks):
        sub = 0
        for row in block:
            ws.append(row)
            r = ws.max_row
            data_rows.append(r)
            if hasattr(row[0], 'year'):
                ws.cell(r, 1).number_format = 'yyyy-mm-dd'
            amt = row[7]
            sub += int(str(amt).replace(',', '')) * (1000 if row[9] == '천원' else 1)
        grand += sub
        ws.append([f'{month_names[mi]} 소계', None, None, None, None, None, None, sub, None, '원', None])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
    ws.append(['합계', None, None, None, None, None, None, grand, None, '원', None])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    hidden = sorted(rng.sample(data_rows, 6))  # 숨은 행: 정상 데이터
    for r in hidden:
        ws.row_dimensions[r].hidden = True
    for col, w in zip('ABCDEFGHIJK', [13, 14, 20, 16, 7, 9, 12, 14, 11, 9, 10]):
        ws.column_dimensions[col].width = w
    wb.save(OUT / '매입내역_정리전.xlsx')
    return len(data_rows), len(hidden), grand


# ---------------------------------------------------------------- 3. 작업일지
REASONS = ['금형 교체', '센서 오류', '자재 대기', '정기점검', '정전', '컨베이어 걸림']
DEFECTS = ['버(날카로운 모서리)', '치수 불량', '도장 긁힘', '용접 불량']
LINES = {'1라인': ('프레스', '브라켓A'), '2라인': ('조립', '커버B')}


def dur_text(rng, minutes):
    if minutes >= 60 and rng.random() < 0.6:
        h, mm = divmod(minutes, 60)
        return f'{h}시간' + (f' {mm}분' if mm else '')
    return f'{minutes}분'


def make_worklog(rng):
    days = []
    d = date(2026, 9, 14)
    while len(days) < 11:
        if d.weekday() < 5 or d == date(2026, 9, 19):  # 평일 + 토요일 1일 특근
            days.append(d)
        d += timedelta(days=1)
    wk = ['월', '화', '수', '목', '금', '토', '일']
    recs, out = [], []
    missing_defect = 0
    for d in days:
        head = rng.choice([f'{d.month}/{d.day}({wk[d.weekday()]})', f'{d.month}월 {d.day}일 {wk[d.weekday()]}요일', f'{d.month:02d}-{d.day:02d}'])
        out.append(f'[{head}]')
        shifts = ['주간'] if d.weekday() == 5 else ['주간', '야간']
        for sh in shifts:
            for ln, (proc, prod) in LINES.items():
                shn = sh + rng.choice(['', '조', '조'])
                if d.weekday() == 5 and ln == '2라인':
                    out.append(f'{ln}은 특근 없이 가동하지 않음.')
                    recs.append((d, ln, sh, prod, 0, 0, '', 0, ''))
                    continue
                total = rng.randint(8, 14) * 100 + rng.choice([0, 20, 50, 80])
                defect = rng.choice([0, 0, 3, 5, 8, 12, 20, 35])
                down, reason = 0, ''
                if rng.random() < 0.45:
                    down, reason = rng.choice([20, 30, 40, 45, 60, 75, 80, 90]), rng.choice(REASONS)
                defect_val = defect
                skip_defect = rng.random() < 0.3 and defect == 0 and missing_defect < 2
                dtype = rng.choice(DEFECTS) if defect else ''
                style = rng.randint(0, 3)
                if style == 0:
                    s = f'{ln}({proc}) {shn} 총 {total:,}개 생산했고 불량 {defect}개 나옴.'
                elif style == 1:
                    s = f'{ln} {shn}: 양품 {total - defect:,}EA, NG {defect}EA. 합계 {total}.'
                elif style == 2:
                    s = f'{shn} {ln} {prod} {total:,}개 완료, 불량은 {defect}개.'
                else:
                    s = f'{ln} {shn} {prod} 생산량 {total}, 불량 {defect}.'
                if skip_defect:  # 불량 언급 없음 -> 정답에서 '확인필요'
                    s = f'{ln}({proc}) {shn} 총 {total:,}개 생산함.'
                    missing_defect += 1
                    defect_val = None
                elif defect:
                    s += f' 불량 유형은 {dtype}.'
                if down:
                    s += rng.choice([f' {reason} 사유로 {dur_text(rng, down)} 멈춤.', f' 다운타임 {dur_text(rng, down)}({reason}).', f' {reason} 때문에 {dur_text(rng, down)} 라인 정지.'])
                out.append(s)
                recs.append((d, ln, sh, prod, total, defect_val, dtype, down, reason))
        out.append('')
    return '\n'.join(out), recs


# ---------------------------------------------------------------- main
def main():
    OUT.mkdir(parents=True, exist_ok=True)
    rng = random.Random(SEED)
    cust = make_customers(rng)
    write_customers(cust)

    blocks, truth = make_purchases(rng)
    n_rows, n_hidden, naive_total = write_purchases(blocks, rng)
    supply_total = sum(t['supply'] for t in truth)
    vat_total = supply_total // 10
    by_item = {}
    month_tot = {}
    for t in truth:
        by_item.setdefault(t['item'], [0, 0])
        by_item[t['item']][0] += 1
        by_item[t['item']][1] += t['supply']
        month_tot[t['date'].month] = month_tot.get(t['date'].month, 0) + t['supply']
    thousand_rows = sum(1 for b in blocks for r in b if r[9] == '천원')
    vat_in_rows = sum(1 for b in blocks for r in b if r[8] == '포함')
    blank_vendor = sum(1 for t in truth if not t['vendor'])

    text, recs = make_worklog(rng)
    (OUT / '작업일지_메모.txt').write_text(text, encoding='utf-8')
    tot_prod = sum(r[4] for r in recs)
    tot_def = sum(r[5] for r in recs if r[5] is not None)
    tot_down = sum(r[7] for r in recs)
    miss_def = sum(1 for r in recs if r[5] is None)
    per_line = {}
    for r in recs:
        p = per_line.setdefault(r[1], [0, 0, 0])
        p[0] += r[4]
        p[1] += r[5] or 0
        p[2] += r[7]

    items_txt = '\n'.join(f'    - {k}: {v[0]}건 / {v[1]:,}원' for k, v in by_item.items())
    line_txt = ', '.join(f'{k} {v[0]:,}개 / {v[1]:,}개 / {v[2]:,}분' for k, v in per_line.items())
    readme = f"""# 실습용 가상 데이터 안내

모든 이름·전화번호·주민등록번호·주소·이메일·차량번호·회사명은 **가상**입니다. 실존 인물·업체와 무관합니다.
이 폴더의 파일과 이 문서는 `python3 samples/make_samples.py` 로 다시 만들 수 있습니다(난수 seed 고정, 실행할 때마다 같은 결과).
아래 "정답"은 생성 스크립트가 계산해서 적은 값입니다. 스킬을 실행한 결과와 맞춰 보세요.

## 1. 고객_방문이력_원본.xlsx (마스킹 실습 → `data-masking`)

- 행 수: {len(cust)}명 (머리글 제외), 시트 `방문이력`
- 열: 고객명, 휴대폰, 주민등록번호, 주소, 이메일, 차량번호, 차종, 최근방문일, 누적금액(원)
- 넣은 함정
  - 휴대폰은 `010-0000-xxxx` 형태, 주민등록번호는 뒤 6자리를 0으로 채운 가짜(실제로 유효하지 않은 형태)
  - 차량번호 3종류가 섞여 있음: `12가3456`, `123가4567`, `서울12가3456`
  - 주소는 `서울시 가상구 샘플로 12`(시·구)와 `경기도 수원시 가상구 ...`(도·시·구)가 섞임. 시·구까지만 남기는 수식은 경기도 주소에서 `경기도 수원시`까지만 남으므로 결과를 눈으로 확인해야 함
  - 동명이인이 일부 있음(이름만으로 사람을 구분하면 안 됨)
- 검산: 마스킹 후에도 행 수는 {len(cust)}행 그대로. 주민등록번호 열은 가리지 말고 삭제하는 것이 정답에 가깝습니다(개인정보위 이용자 가이드: 주민등록번호는 아예 입력하지 않기).
- 실습 방법: 이 파일을 AI에 올리기 전에, 열 이름과 2~3행만 AI에 알려 주고 엑셀 수식을 받아 PC에서 가린 뒤 올려 보세요.

## 2. 매입내역_정리전.xlsx (엑셀 정리 실습 → `excel-cleanup`)

- 시트 `매입내역`, 기간 2026-07 ~ 2026-09, 데이터 행 {n_rows}행
- 넣은 함정
  - 1행 제목 병합 셀(A1:K1), 2~3행 머리글이 2줄(그룹 머리글 병합 + 세부 머리글)
  - 월별 `소계` 행 3개와 맨 아래 `합계` 행(데이터가 아님. 이 행의 숫자는 부가세 포함 금액이 섞인 채 더한 값이라 정답과 다름)
  - 금액 단위 혼재: `금액단위` 열이 `천원`인 행 {thousand_rows}개(단가·금액 모두 천원 기준)
  - 부가세 혼재: `부가세구분` 열이 `포함`인 행 {vat_in_rows}개(단가·금액이 부가세 포함가. 공급가액 = 금액 ÷ 1.1)
  - 문자로 저장된 숫자: `1,200`, `300개`, `₩12,000`, `12,000원`, 앞뒤 공백
  - 날짜 형식 혼재: 진짜 날짜, `2026.07.05`, `26/7/5`, `20260705`, 엑셀 일련번호(숫자만), `7월 5일`(연도 없음 → 제목의 2026년 3분기로 판단)
  - 같은 품목의 다른 표기: `볼트M8` / `볼트 M8` / `BOLT-M8` / 전각 `볼트 Ｍ８` 등, 품목명 앞뒤 공백, 전각 공백
  - 숨은 행 {n_hidden}개(정상 데이터이므로 정리 후에도 포함해야 함)
  - 거래처 빈칸 {blank_vendor}행(임의 값으로 채우지 말고 빈칸 + 확인필요로 표시)
- 정답 (정리 후)
  - 데이터 행 수: **{len(truth)}행** (소계 3개, 합계 1개, 제목·머리글 제외)
  - 공급가액(부가세 별도, 원 단위) 합계: **{supply_total:,}원**, 부가세 10%: {vat_total:,}원
  - 월별 공급가액: 7월 {month_tot[7]:,}원 / 8월 {month_tot[8]:,}원 / 9월 {month_tot[9]:,}원
  - 파일 맨 아래 합계 행의 값(틀린 값): {naive_total:,}원 → 정리 후 합계와 다른 것이 정상
  - 품목별 (건수 / 공급가액):
{items_txt}

## 3. 작업일지_메모.txt (메모를 표로 → `text-to-table`)

- 2026-09-14 ~ 2026-09-25 평일 10일 + 2026-09-19(토) 특근 1일, 라인 2개(1라인 프레스·브라켓A, 2라인 조립·커버B), 주간/야간 교대
- 정답 기준: **생산수량 = 불량 포함 총 생산 수량**. `양품 ... NG ...` 형태로 적힌 줄은 양품과 NG를 더하지 않고, 줄 끝에 적힌 합계(총 생산)를 그대로 쓴다.
- 넣은 함정
  - 날짜 머리 표기가 `9/14(월)`, `9월 15일 화요일`, `09-16` 등으로 섞임
  - 주간/주간조, 야간/야간조 표기 혼재, `불량`과 `NG` 혼용, `다운타임`·`정지`·`멈춤` 혼용
  - 비가동 시간이 `1시간 20분`, `75분` 등으로 섞임
  - 불량 수량을 말하지 않은 교대 {miss_def}건(0으로 채우지 말고 `확인필요`)
  - 토요일 2라인은 가동하지 않음(생산 0이 정상)
- 정답
  - 교대·라인 기록 수: **{len(recs)}건**
  - 총 생산수량: **{tot_prod:,}개**
  - 총 불량수량: **{tot_def:,}개** (불량 수량이 적히지 않은 {miss_def}건 제외)
  - 총 비가동시간: **{tot_down:,}분**
  - 라인별 (생산 / 불량 / 비가동): {line_txt}
"""
    (OUT / 'README.md').write_text(readme, encoding='utf-8')
    print(f'customers={len(cust)} purchase_rows={len(truth)} supply={supply_total} worklog_recs={len(recs)} prod={tot_prod} def={tot_def} down={tot_down} miss_def={miss_def}')


if __name__ == '__main__':
    main()
